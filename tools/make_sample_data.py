"""Generate the sample LCNRV data workbook (data/LCNRV_Data.xlsx).

The workbook holds three Excel tables that the Power BI model reads:
  tblProducts  - SKU master
  tblInventory - month-end holdings: quantity and unit cost (COGS per unit)
  tblPrices    - month-end expected selling price and costs to complete / sell

Replace the sample rows with your own data, keeping the table and column names.
"""
import datetime as dt
import random
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

random.seed(42)

PRODUCTS = [
    # SKU, description, category, base unit cost, base selling price, price trend / month
    ("SKU-1001", "Aluminium Sheet 2mm", "Raw Materials", 18.50, 26.00, -0.010),
    ("SKU-1002", "Copper Wire 10m", "Raw Materials", 42.00, 55.00, -0.035),
    ("SKU-1003", "Steel Bracket", "Components", 3.20, 4.10, -0.020),
    ("SKU-1004", "Control Board v2", "Components", 64.00, 92.00, -0.060),
    ("SKU-1005", "Sensor Module", "Components", 21.00, 34.00, 0.005),
    ("SKU-2001", "Smart Thermostat", "Finished Goods", 85.00, 129.00, -0.015),
    ("SKU-2002", "Smart Thermostat (Legacy)", "Finished Goods", 80.00, 99.00, -0.045),
    ("SKU-2003", "Wireless Hub", "Finished Goods", 47.00, 79.00, 0.000),
    ("SKU-2004", "Door Sensor 3-pack", "Finished Goods", 22.00, 30.00, -0.030),
    ("SKU-2005", "Seasonal Gift Bundle", "Finished Goods", 35.00, 60.00, None),  # seasonal
]
LOCATIONS = ["Main Warehouse", "East DC"]
CATEGORY_COSTS = {  # costs to complete per unit, selling cost % of price
    "Raw Materials": (2.00, 0.04),
    "Components": (1.50, 0.06),
    "Finished Goods": (0.00, 0.12),
}


def month_ends(start_year, start_month, n):
    out, y, m = [], start_year, start_month
    for _ in range(n):
        nxt = dt.date(y + (m == 12), m % 12 + 1, 1)
        out.append(nxt - dt.timedelta(days=1))
        y, m = nxt.year, nxt.month
    return out


def main(path):
    months = month_ends(2025, 10, 12)  # Oct-2025 .. Sep-2026
    inv_rows, price_rows = [], []
    for sku, _, cat, cost, price, trend in PRODUCTS:
        ctc, sell_pct = CATEGORY_COSTS[cat]
        for i, me in enumerate(months):
            # Unit cost drifts slightly (weighted-average cost moves with purchases)
            unit_cost = round(cost * (1 + 0.004 * i) * random.uniform(0.99, 1.01), 2)
            if trend is None:  # seasonal bundle: full price Oct-Dec, clearance Jan-Apr, recovers after
                factor = 1.0 if me.month in (10, 11, 12) else (0.55 if me.month <= 4 else 0.85)
            else:
                factor = (1 + trend) ** i
            sp = round(price * factor * random.uniform(0.98, 1.02), 2)
            price_rows.append((me, sku, sp, ctc, sell_pct))
            for loc in LOCATIONS:
                if sku == "SKU-2002" and loc == "East DC" and i >= 8:
                    continue  # legacy item cleared out of East DC
                qty = max(0, int(random.gauss(900 if cat != "Finished Goods" else 400, 120) * (0.6 if loc == "East DC" else 1)))
                inv_rows.append((me, sku, loc, qty, unit_cost))
    # Leave one price missing so the "missing price" check has something to show
    price_rows = [r for r in price_rows if not (r[0] == months[-1] and r[1] == "SKU-1005")]

    wb = Workbook()
    sheets = [
        ("Products", "tblProducts", ["SKU", "Description", "Category"],
         [(p[0], p[1], p[2]) for p in PRODUCTS], {}),
        ("Inventory", "tblInventory", ["MonthEnd", "SKU", "Location", "Quantity", "UnitCost"],
         inv_rows, {"A": "yyyy-mm-dd", "D": "#,##0", "E": "#,##0.00"}),
        ("Prices", "tblPrices", ["MonthEnd", "SKU", "SellingPrice", "CostToComplete", "SellingCostPct"],
         price_rows, {"A": "yyyy-mm-dd", "C": "#,##0.00", "D": "#,##0.00", "E": "0.0%"}),
    ]
    for idx, (title, tname, headers, rows, fmts) in enumerate(sheets):
        ws = wb.active if idx == 0 else wb.create_sheet()
        ws.title = title
        ws.append(headers)
        for r in rows:
            ws.append(list(r))
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for col, fmt in fmts.items():
            for c in ws[col][1:]:
                c.number_format = fmt
        for i, h in enumerate(headers):
            ws.column_dimensions[chr(65 + i)].width = max(14, len(h) + 4)
        ref = f"A1:{chr(64 + len(headers))}{len(rows) + 1}"
        t = Table(displayName=tname, ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(t)
        ws.freeze_panes = "A2"
    ws = wb.create_sheet("ReadMe", 0)
    notes = [
        "LCNRV data template for the Power BI model",
        "",
        "tblInventory: one row per month-end x SKU x location. Quantity = units on hand; UnitCost = cost per unit (COGS / weighted-average cost).",
        "tblPrices: one row per month-end x SKU. NRV per unit = SellingPrice x (1 - SellingCostPct) - CostToComplete.",
        "tblProducts: one row per SKU.",
        "MonthEnd can be any date in the month; the model snaps it to the month end.",
        "Keep the table names and column headers; add or replace rows freely.",
    ]
    for n in notes:
        ws.append([n])
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 120
    wb.save(path)
    print(f"wrote {path}: {len(inv_rows)} inventory rows, {len(price_rows)} price rows")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "data/LCNRV_Data.xlsx")
